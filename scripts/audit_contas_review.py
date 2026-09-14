from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


FILE = Path(r"C:\dev\hpo-ptbr-lab\output\Contabilidade_Geral_e_Avancada_Revisao_Cumulativa.docx")
doc = Document(FILE)
text = "\n".join(p.text for p in doc.paragraphs)

required = [
    "Aula 00",
    "Aula 01",
    "Aula 02 — Contas",
    "Conceito e finalidade das contas",
    "Partidas dobradas e natureza dos saldos",
    "Plano de contas",
    "Teorias das contas",
    "Contas sintéticas, analíticas e de compensação",
    "Mapa essencial das principais contas",
    "Revisão rápida da Aula 02",
    "páginas 3 a 26",
    "página 27 inicia",
]
missing = [item for item in required if item not in text]
assert not missing, f"Missing required content: {missing}"

for forbidden in ["==43512d==", "C:\\Users\\xboxf", "@gmail.com"]:
    assert forbidden not in text, f"Forbidden marker found: {forbidden}"

assert text.count("Aula 02 — Contas") >= 2, "Aula 02 must appear in TOC and heading"
assert "Fechamento da Aula 01" in text, "Existing Aula 01 close was lost"
assert "Fechamento da Aula 02" in text, "Aula 02 close missing"

with ZipFile(FILE) as zf:
    xml = zf.read("word/document.xml")

from lxml import etree

root = etree.fromstring(xml)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
bookmarks = root.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
anchors = root.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
assert len(bookmarks) == len(set(bookmarks)), "Duplicate bookmark names"
missing_anchors = sorted(set(anchors) - set(bookmarks))
assert not missing_anchors, f"TOC anchors without bookmarks: {missing_anchors}"

new_bookmarks = {
    "aula02_contas",
    "aula02_conceito",
    "aula02_partidas",
    "aula02_razo",
    "aula02_plano",
    "aula02_teorias",
    "aula02_classificacao",
    "aula02_mapa",
    "aula02_revisao",
}
assert new_bookmarks <= set(bookmarks), f"Missing new bookmarks: {new_bookmarks - set(bookmarks)}"
assert new_bookmarks <= set(anchors), f"Missing new links: {new_bookmarks - set(anchors)}"

for idx, table in enumerate(doc.tables, start=1):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    assert tr_pr.find(qn("w:tblHeader")) is not None, f"Table {idx} lacks repeating header"

heading_text = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
idx_00 = next(i for i, value in enumerate(heading_text) if value.startswith("Aula 00"))
idx_01 = next(i for i, value in enumerate(heading_text) if value.startswith("Aula 01"))
idx_02 = next(i for i, value in enumerate(heading_text) if value.startswith("Aula 02"))
assert idx_00 < idx_01 < idx_02, "Aula heading order is incorrect"

print("CUSTOM AUDIT: OK")
print(f"Bookmarks: {len(bookmarks)} | Internal links: {len(anchors)}")
print(f"Paragraphs: {len(doc.paragraphs)} | Tables: {len(doc.tables)}")
print(f"File size: {FILE.stat().st_size} bytes")
