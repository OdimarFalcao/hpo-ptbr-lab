from __future__ import annotations

import copy
import os
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\dev\hpo-ptbr-lab")
SOURCE = ROOT / "tmp" / "contas_merge_20260901" / "Contabilidade_Geral_e_Avancada_Revisao_Cumulativa.docx"
OUTPUT = ROOT / "output" / "Contabilidade_Geral_e_Avancada_Revisao_Cumulativa.docx"

PAGE_WIDTH_DXA = 9288
TABLE_INDENT_DXA = 110
GREEN = "38761D"
LIGHT_GREEN = "D9EAD3"
PALE_GREEN = "EEF6EA"
GRID = "B7B7B7"


def set_text_preserving_first_run(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_page_break_before(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    el = p_pr.find(qn("w:pageBreakBefore"))
    if el is None:
        el = OxmlElement("w:pageBreakBefore")
        p_pr.append(el)


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_bullet(paragraph, level: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)


def add_bullet(doc: Document, text: str, bold_prefix: str | None = None, level: int = 0):
    p = doc.add_paragraph(style="normal")
    set_bullet(p, level)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_labeled_paragraph(doc: Document, label: str, text: str):
    p = doc.add_paragraph(style="normal")
    p.add_run(label).bold = True
    p.add_run(text)
    return p


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def set_table_properties(table, widths: list[int]) -> None:
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout", "tblBorders"):
        old = tbl_pr.find(qn(f"w:{tag}"))
        if old is not None:
            tbl_pr.remove(old)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), GRID)
        borders.append(border)
    tbl_pr.extend([tbl_w, tbl_ind, tbl_layout, borders])

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], ratios: list[float]):
    widths = [int(PAGE_WIDTH_DXA * value / sum(ratios)) for value in ratios]
    widths[-1] += PAGE_WIDTH_DXA - sum(widths)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.text = value
        set_cell_shading(cell, LIGHT_GREEN)
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(GREEN)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            row.cells[idx].text = value
            if row_idx % 2:
                set_cell_shading(row.cells[idx], PALE_GREEN)
    set_table_properties(table, widths)
    return table


def normalize_existing_tables(doc: Document) -> None:
    for table in doc.tables:
        col_count = len(table.columns)
        existing = []
        for grid_col in table._tbl.tblGrid.findall(qn("w:gridCol"))[:col_count]:
            raw = grid_col.get(qn("w:w"))
            try:
                existing.append(int(float(raw)))
            except (TypeError, ValueError):
                existing.append(1)
        if len(existing) != col_count or sum(existing) <= 0:
            existing = [1] * col_count
        widths = [int(PAGE_WIDTH_DXA * value / sum(existing)) for value in existing]
        widths[-1] += PAGE_WIDTH_DXA - sum(widths)
        set_table_properties(table, widths)
        set_repeat_table_header(table.rows[0])


def add_heading(doc: Document, text: str, level: int, bookmark: tuple[str, int] | None = None, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    if page_break:
        set_page_break_before(p)
    if bookmark:
        add_bookmark(p, bookmark[0], bookmark[1])
    return p


def find_paragraph(doc: Document, needle: str):
    for p in doc.paragraphs:
        if needle.casefold() in p.text.casefold():
            return p
    raise ValueError(f"Paragraph not found: {needle}")


def insert_toc_entries(doc: Document, source_note, entries: list[tuple[str, str]]) -> None:
    peer = find_paragraph(doc, "Revisão rápida da Aula 01")
    for text, anchor in entries:
        new_p = copy.deepcopy(peer._p)
        for child in list(new_p):
            if child.tag != qn("w:pPr"):
                new_p.remove(child)
        temp_paragraph = type(peer)(new_p, peer._parent)
        add_internal_link(temp_paragraph, text, anchor)
        source_note._p.addprevious(new_p)


def append_aula_02(doc: Document) -> None:
    add_heading(doc, "Aula 02 — Contas", 1, ("aula02_contas", 100), page_break=True)
    p = doc.add_paragraph(style="normal")
    p.add_run("Objetivo da etapa: ").bold = True
    p.add_run(
        "dominar o conceito de conta, a lógica de débito e crédito, a estrutura do plano de contas, "
        "as teorias classificatórias e as contas que mais geram confusão em provas fiscais."
    )

    add_heading(doc, "Conceito e finalidade das contas", 2, ("aula02_conceito", 101))
    doc.add_paragraph(
        "Conta contábil é o título que identifica um elemento patrimonial ou de resultado e permite "
        "registrar, acumular e controlar suas variações. Seu saldo traduz, em dado momento, a posição "
        "quantitativa daquele elemento.",
        style="normal",
    )
    add_labeled_paragraph(
        doc,
        "Lei nº 6.404/1976, art. 176, § 2º: ",
        "nas demonstrações, contas semelhantes podem ser agrupadas; pequenos saldos podem ser reunidos, "
        "desde que sua natureza seja indicada e o total não ultrapasse 10% do grupo. Designações genéricas "
        "como “diversas contas” e “contas-correntes” não são admitidas.",
    )
    add_bullet(doc, "Conta patrimonial: representa ativo, passivo ou patrimônio líquido.", "Conta patrimonial:")
    add_bullet(doc, "Conta de resultado: representa receita ou despesa e é encerrada na apuração do resultado.", "Conta de resultado:")
    add_bullet(doc, "Conta de compensação: controla ato administrativo relevante que ainda não altera o patrimônio.", "Conta de compensação:")

    add_heading(doc, "Partidas dobradas e natureza dos saldos", 2, ("aula02_partidas", 102))
    doc.add_paragraph(
        "Pelo método das partidas dobradas, todo lançamento envolve ao menos um débito e um crédito, "
        "sempre com igualdade entre os valores debitados e creditados. Débito e crédito são convenções "
        "técnicas: não significam, por si, algo favorável ou desfavorável.",
        style="normal",
    )
    add_table(
        doc,
        ["Grupo", "Natureza normal", "Aumenta por", "Diminui por"],
        [
            ["Ativo", "Devedora", "Débito", "Crédito"],
            ["Despesa", "Devedora", "Débito", "Crédito"],
            ["Passivo", "Credora", "Crédito", "Débito"],
            ["Patrimônio líquido", "Credora", "Crédito", "Débito"],
            ["Receita", "Credora", "Crédito", "Débito"],
            ["Retificadora do ativo", "Credora", "Crédito", "Débito"],
            ["Retificadora do passivo ou do PL", "Devedora", "Débito", "Crédito"],
        ],
        [2.3, 1.7, 1.5, 1.5],
    )
    add_labeled_paragraph(
        doc,
        "Extrato bancário: ",
        "o banco registra a conta sob sua própria perspectiva. O depósito do cliente é obrigação do banco; "
        "por isso aparece como crédito no extrato, embora “Bancos” seja conta devedora para a empresa.",
    )
    add_labeled_paragraph(
        doc,
        "Contas estáveis e instáveis: ",
        "estável é a conta que normalmente conserva um único tipo de saldo; instável pode apresentar saldo "
        "devedor ou credor conforme os fatos registrados.",
    )

    add_heading(doc, "Função, estrutura e razonete", 2, ("aula02_razo", 103))
    add_table(
        doc,
        ["Elemento", "Leitura para prova"],
        [
            ["Função", "Explica o que a conta registra, controla e evidencia."],
            ["Estrutura", "Reúne título, débitos, créditos, histórico, data e saldo."],
            ["Razonete", "Representação em T: débitos à esquerda e créditos à direita."],
            ["Saldo", "Diferença entre o total de débitos e o total de créditos."],
        ],
        [2.2, 5.8],
    )
    add_bullet(doc, "Total de débitos maior que o de créditos → saldo devedor.")
    add_bullet(doc, "Total de créditos maior que o de débitos → saldo credor.")
    add_bullet(doc, "Totais iguais → conta saldada, com saldo zero.")

    add_heading(doc, "Plano de contas", 2, ("aula02_plano", 104))
    doc.add_paragraph(
        "Plano de contas é o conjunto organizado de contas utilizado pela entidade. Ele deve ser coerente "
        "com a atividade, o porte, as necessidades gerenciais e as exigências legais; por isso, não existe "
        "um plano universal e imutável para todas as empresas.",
        style="normal",
    )
    add_table(
        doc,
        ["Componente", "Finalidade"],
        [
            ["Elenco de contas", "Relação ordenada e codificada das contas."],
            ["Manual de contas", "Descreve função, funcionamento e critérios de cada conta."],
            ["Modelos padronizados", "Orientam a elaboração das demonstrações e outros relatórios."],
            ["Hierarquia", "Organiza grupos, subgrupos e contas em níveis sintéticos e analíticos."],
        ],
        [2.5, 5.5],
    )

    add_heading(doc, "Teorias das contas", 2, ("aula02_teorias", 105))
    add_table(
        doc,
        ["Teoria", "Classes de contas"],
        [
            ["Personalista", "Agentes consignatários: bens; agentes correspondentes: direitos e obrigações; proprietários: PL, receitas e despesas."],
            ["Materialista", "Contas integrais: bens, direitos e obrigações; contas diferenciais: PL, receitas e despesas."],
            ["Patrimonialista", "Contas patrimoniais: ativo, passivo e PL; contas de resultado: receitas e despesas. É a teoria atualmente adotada."],
        ],
        [2.0, 6.0],
    )
    add_labeled_paragraph(doc, "Mnemônico de prova: ", "na teoria materialista, INTEGRAIS = bens, direitos e obrigações; DIFERENCIAIS = PL, receitas e despesas.")

    add_heading(doc, "Contas sintéticas, analíticas e de compensação", 2, ("aula02_classificacao", 106))
    add_bullet(doc, "Sintética: agrega outras contas e normalmente não recebe lançamento direto; exemplo: Bancos.", "Sintética:")
    add_bullet(doc, "Analítica: apresenta o maior grau de detalhamento e recebe os lançamentos; exemplo: Banco Alfa.", "Analítica:")
    add_bullet(doc, "Compensação ou extrapatrimonial: registra controles de atos relevantes, sem modificar imediatamente ativo, passivo ou PL.", "Compensação ou extrapatrimonial:")
    doc.add_paragraph(
        "A classificação sintética ou analítica depende do nível de detalhamento adotado no plano de contas: "
        "uma mesma denominação pode ser sintética em uma estrutura e analítica em outra.",
        style="normal",
    )

    add_heading(doc, "Mapa essencial das principais contas", 2, ("aula02_mapa", 107))
    add_table(
        doc,
        ["Conta ou família", "Classificação / natureza", "Ponto de atenção"],
        [
            ["Caixa, Bancos, Clientes, Estoques", "Ativo / devedora", "Aplicações de recursos; aumentam por débito."],
            ["Depreciação e perdas estimadas", "Retificadoras do ativo / credora", "Reduzem o saldo contábil do ativo relacionado."],
            ["Fornecedores, tributos a recolher, empréstimos", "Passivo / credora", "Obrigações; aumentam por crédito."],
            ["Capital a integralizar", "Retificadora do PL / devedora", "Parcela subscrita ainda não entregue pelos sócios."],
            ["Ações em tesouraria", "Retificadora do PL / devedora", "Ações próprias adquiridas pela companhia."],
            ["Adiantamento a fornecedor", "Ativo / devedora", "Direito de receber bem ou serviço."],
            ["Adiantamento de cliente", "Passivo / credora", "Obrigação de entregar bem ou prestar serviço."],
            ["Despesa antecipada", "Ativo / devedora", "Permanece no ativo enquanto o benefício não for consumido."],
            ["Dividendos adicionais propostos", "PL / credora enquanto não constituírem obrigação", "Não confundir com dividendos obrigatórios a pagar."],
            ["Ajustes de avaliação patrimonial", "PL / devedora ou credora", "O sentido depende da variação reconhecida."],
            ["Ajustes de exercícios anteriores", "PL / devedora ou credora", "Podem aumentar ou reduzir lucros acumulados."],
        ],
        [2.6, 2.4, 4.0],
    )
    add_labeled_paragraph(
        doc,
        "Como estudar o anexo do PDF: ",
        "use o elenco de contas como material de consulta e para reconhecer famílias e naturezas; não é necessário memorizar imediatamente todas as rubricas.",
    )

    add_heading(doc, "Palavras-chave e pegadinhas de prova", 2)
    for text in [
        "partidas dobradas → soma dos débitos sempre igual à soma dos créditos;",
        "ativo e despesa → natureza devedora; passivo, PL e receita → natureza credora;",
        "retificadora → natureza oposta à do grupo que reduz;",
        "débito e crédito → convenções, não sinônimos de perda e ganho;",
        "plano de contas → adaptável à entidade, e não modelo único obrigatório;",
        "conta analítica → recebe lançamento; conta sintética → agrega detalhes;",
        "adiantamento → classifique conforme quem tem o direito e quem possui a obrigação;",
        "extrato bancário → está sob a ótica do banco.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Revisão rápida da Aula 02", 1, ("aula02_revisao", 108), page_break=True)
    add_heading(doc, "Checklist de 3 minutos", 2)
    for text in [
        "Sei explicar por que todo lançamento mantém débitos e créditos em igualdade?",
        "Reconheço a natureza normal de ativo, passivo, PL, receitas e despesas?",
        "Consigo inverter a natureza quando a conta é retificadora?",
        "Diferencio função, estrutura, razonete e saldo?",
        "Memorizei as três teorias e suas classes de contas?",
        "Diferencio conta sintética, analítica e de compensação?",
        "Classifico corretamente adiantamentos e despesas antecipadas?",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Autoavaliação expressa", 2)
    questions = [
        "1. Por que o crédito do extrato bancário não transforma “Bancos” em conta credora para a empresa?",
        "2. Qual é a natureza de uma conta retificadora do ativo?",
        "3. Como a teoria materialista separa contas integrais e diferenciais?",
        "4. Qual conta recebe diretamente os lançamentos: sintética ou analítica?",
        "5. Adiantamento de cliente representa direito ou obrigação da entidade?",
    ]
    for question in questions:
        doc.add_paragraph(question, style="normal")

    p = doc.add_paragraph(style="normal")
    p.add_run("Gabarito mental: ").bold = True
    p.add_run(
        "1) porque o extrato segue a ótica do banco; 2) credora; 3) integrais = bens, direitos e obrigações, "
        "diferenciais = PL, receitas e despesas; 4) analítica; 5) obrigação, portanto passivo."
    )
    p = doc.add_paragraph(style="normal")
    p.add_run("Fechamento da Aula 02: ").bold = True
    p.add_run(
        "conteúdo teórico das páginas 3 a 26 consolidado. A página 27 inicia “Questões Comentadas — Contas — Cebraspe” e não foi incorporada."
    )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    if any("Aula 02 — Contas" in p.text for p in doc.paragraphs):
        raise RuntimeError("The source already contains Aula 02; refusing to duplicate it.")

    subtitle = find_paragraph(doc, "Concursos fiscais | Aulas 00 e 01")
    set_text_preserving_first_run(
        subtitle,
        "Concursos fiscais | Aulas 00 a 02 — Aspectos introdutórios, patrimônio e contas",
    )

    source_note = find_paragraph(doc, "Fonte desta etapa:")
    source_text = (
        source_note.text.rstrip()
        + " Aula 02: 02-Contas.pdf, páginas 3 a 26; a página 27 inicia Questões Comentadas — Contas — Cebraspe e não foi incorporada."
    )
    set_text_preserving_first_run(source_note, source_text)

    insert_toc_entries(
        doc,
        source_note,
        [
            ("Aula 02 — Contas", "aula02_contas"),
            ("Conceito e finalidade das contas", "aula02_conceito"),
            ("Partidas dobradas e natureza dos saldos", "aula02_partidas"),
            ("Função, estrutura e razonete", "aula02_razo"),
            ("Plano de contas", "aula02_plano"),
            ("Teorias das contas", "aula02_teorias"),
            ("Contas sintéticas, analíticas e de compensação", "aula02_classificacao"),
            ("Mapa essencial das principais contas", "aula02_mapa"),
            ("Revisão rápida da Aula 02", "aula02_revisao"),
        ],
    )

    append_aula_02(doc)
    normalize_existing_tables(doc)

    doc.core_properties.subject = "Revisão cumulativa para concursos fiscais — Contabilidade, Aulas 00 a 02"
    doc.core_properties.comments = (
        "Aula 02 — Contas incorporada com base nas páginas 3 a 26 de 02-Contas.pdf; "
        "questões comentadas, iniciadas na página 27, foram excluídas."
    )
    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
