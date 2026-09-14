from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\dev\hpo-ptbr-lab")
SOURCE = ROOT / "tmp" / "aula04_20260911" / "revisao_atual.docx"
OUTPUT = ROOT / "output" / "Contabilidade - Revisão Cumulativa - Contas, Atos e Fatos Contábeis.docx"


def find_paragraph(
    doc: Document,
    exact: str | None = None,
    startswith: str | None = None,
    style_name: str | None = None,
):
    for paragraph in doc.paragraphs:
        if style_name is not None and paragraph.style.name != style_name:
            continue
        if exact is not None and paragraph.text == exact:
            return paragraph
        if startswith is not None and paragraph.text.startswith(startswith):
            return paragraph
    raise ValueError(
        f"Paragraph not found: exact={exact!r} startswith={startswith!r} style={style_name!r}"
    )


def replace_text_keep_format(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def clone_toc_paragraph(peer, text: str, anchor: str):
    new_p = copy.deepcopy(peer._p)
    hyperlinks = new_p.xpath(".//w:hyperlink")
    if not hyperlinks:
        raise RuntimeError("TOC peer has no internal hyperlink")
    hyperlinks[0].set(qn("w:anchor"), anchor)
    text_nodes = new_p.xpath(".//w:t")
    if not text_nodes:
        raise RuntimeError("TOC peer has no visible text")
    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ""
    return new_p


def replace_hyperlink_text(paragraph, text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:hyperlink//w:t")
    if not text_nodes:
        raise RuntimeError(f"Paragraph has no hyperlink text: {paragraph.text}")
    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ""


def clone_num_pr(peer, paragraph) -> None:
    peer_ppr = peer._p.pPr
    if peer_ppr is None or peer_ppr.numPr is None:
        raise RuntimeError(f"Bullet peer has no numbering: {peer.text}")
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.numPr is not None:
        p_pr.remove(p_pr.numPr)
    p_pr.append(copy.deepcopy(peer_ppr.numPr))


def add_bullet(doc: Document, peer, text: str, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph(style=peer.style)
    clone_num_pr(peer, paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_heading(doc: Document, text: str, level: int, bookmark: tuple[str, int] | None = None, page_break: bool = False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    if bookmark:
        add_bookmark(paragraph, bookmark[0], bookmark[1])
    return paragraph


def add_labeled_paragraph(doc: Document, label: str, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)
    return paragraph


def replace_element(parent, old, new) -> None:
    index = parent.index(old)
    parent.remove(old)
    parent.insert(index, new)


def clone_run_properties(peer_run, run) -> None:
    if peer_run is None or peer_run._r.rPr is None:
        return
    if run._r.rPr is not None:
        run._r.remove(run._r.rPr)
    run._r.insert(0, copy.deepcopy(peer_run._r.rPr))


def replace_cell_text(cell, text: str, peer_cell) -> None:
    peer_paragraph = peer_cell.paragraphs[0]
    peer_run = peer_paragraph.runs[0] if peer_paragraph.runs else None
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = peer_paragraph.style
    if peer_paragraph._p.pPr is not None:
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        paragraph._p.insert(0, copy.deepcopy(peer_paragraph._p.pPr))
    run = paragraph.add_run(text)
    clone_run_properties(peer_run, run)


def style_table_from_peer(table, peer_table) -> None:
    replace_element(table._tbl, table._tbl.tblPr, copy.deepcopy(peer_table._tbl.tblPr))
    replace_element(table._tbl, table._tbl.tblGrid, copy.deepcopy(peer_table._tbl.tblGrid))
    for row_index, row in enumerate(table.rows):
        peer_row = peer_table.rows[0] if row_index == 0 else peer_table.rows[min(1, len(peer_table.rows) - 1)]
        for col_index, cell in enumerate(row.cells):
            peer_cell = peer_row.cells[min(col_index, len(peer_row.cells) - 1)]
            if cell._tc.tcPr is not None:
                cell._tc.remove(cell._tc.tcPr)
            if peer_cell._tc.tcPr is not None:
                cell._tc.insert(0, copy.deepcopy(peer_cell._tc.tcPr))
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def add_two_col_table(doc: Document, peer_table, headers: tuple[str, str], rows: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=2)
    for col, header in enumerate(headers):
        replace_cell_text(table.rows[0].cells[col], header, peer_table.rows[0].cells[col])
    for values in rows:
        row = table.add_row()
        for col, value in enumerate(values):
            peer_cell = peer_table.rows[min(1, len(peer_table.rows) - 1)].cells[col]
            replace_cell_text(row.cells[col], value, peer_cell)
    style_table_from_peer(table, peer_table)
    doc.add_paragraph("")
    return table


def move_new_block_before(doc: Document, first_element, target_element) -> None:
    body = doc.element.body
    elements = list(body)
    start = elements.index(first_element)
    target = elements.index(target_element)
    if start >= target:
        moving = [element for element in elements[start:] if element.tag != qn("w:sectPr")]
        for element in moving:
            target_element.addprevious(element)


def insert_new_paragraphs_before(target, paragraphs) -> None:
    for paragraph in paragraphs:
        target._p.addprevious(paragraph._p)


def append_aula04(doc: Document, bullet_peer, table_peer) -> None:
    first = add_heading(doc, "Aula 04 - Escrituração", 1, ("aula_04", 100), page_break=True)
    add_labeled_paragraph(
        doc,
        "Objetivo da etapa: ",
        "compreender como os fatos contábeis são registrados, corrigidos, transferidos aos livros e conferidos por meio do balancete de verificação.",
    )

    add_heading(doc, "Aspectos gerais e obrigatoriedade", 2, ("escrituracao_aspectos", 101))
    doc.add_paragraph(
        "Escrituração contábil é a técnica de registrar os fatos contábeis por lançamentos, em livros próprios, pelo método das partidas dobradas. Os registros devem ser permanentes, uniformes, apoiados em documentação e compatíveis com a legislação comercial e os princípios contábeis.",
        style="Normal",
    )
    add_labeled_paragraph(
        doc,
        "Legislação tributária ou especial: ",
        "quando exigir critérios ou demonstrativos diferentes, a companhia deve atendê-los em livros ou registros auxiliares, sem alterar a escrituração mercantil e as demonstrações regidas pela Lei 6.404/1976.",
    )
    add_labeled_paragraph(
        doc,
        "Obrigatoriedade: ",
        "o Código Civil exige sistema de contabilidade uniforme para o empresário e a sociedade empresária. A aula destaca as dispensas legais do produtor rural e do pequeno empresário, observado o enquadramento aplicável.",
    )

    add_heading(doc, "Livros de escrituração", 2, ("livros_escrituracao", 102))
    add_two_col_table(
        doc,
        table_peer,
        ("Livro", "Características essenciais"),
        [
            ("Diário", "Obrigatório, principal e cronológico. Registra todos os fatos em ordem temporal e está sujeito às formalidades intrínsecas e extrínsecas. Não pode ser substituído pelo Razão."),
            ("Razão", "Principal e sistemático. Controla separadamente o movimento de cada conta. É facultativo pela legislação societária e obrigatório, segundo o RIR/2018, para entidades tributadas pelo lucro real."),
            ("Caixa", "Registra cronologicamente recebimentos e pagamentos. Não substitui o Diário nem o Razão."),
            ("Livros das sociedades anônimas", "Além dos livros gerais, incluem registros e transferências de ações e partes beneficiárias, atas, presença de acionistas e pareceres do conselho fiscal, conforme o art. 100 da Lei 6.404/1976."),
        ],
    )
    add_labeled_paragraph(doc, "Razonete: ", "é a representação simplificada do Razão em forma de T, com débitos à esquerda e créditos à direita.")

    add_heading(doc, "Formalidades da escrituração", 2, ("formalidades_escrituracao", 103))
    doc.add_paragraph(
        "Segundo a ITG 2000, o registro deve conter data, conta debitada, conta creditada, histórico que expresse a essência econômica, valor e identificação unívoca dos registros que formam o mesmo lançamento.",
        style="Normal",
    )
    add_labeled_paragraph(
        doc,
        "Escrituração resumida: ",
        "o Diário pode receber totais de operações numerosas por período que não exceda trinta dias, desde que haja escrituração analítica em livros auxiliares regularmente autenticados e documentação que permita a verificação.",
    )
    add_two_col_table(
        doc,
        table_peer,
        ("Formalidade", "Como reconhecer"),
        [
            ("Extrínseca", "Relaciona-se à apresentação material do livro e dificulta adulterações: encadernação, folhas numeradas, termos de abertura e encerramento e assinaturas; no meio digital, assinaturas digitais e autenticação quando exigida. A inobservância invalida o livro."),
            ("Intrínseca", "Relaciona-se à fidedignidade do lançamento: idioma e moeda nacionais, forma contábil, ordem cronológica, ausência de espaços, borrões, rasuras ou emendas e suporte documental. A inobservância invalida apenas o registro atingido."),
        ],
    )

    add_heading(doc, "Lançamentos contábeis e fórmulas", 2, ("lancamentos_formulas", 104))
    doc.add_paragraph(
        "No Livro Diário, o lançamento completo reúne local e data, conta ou contas debitadas, conta ou contas creditadas, histórico e valor. Na forma manual, a preposição “a” indica a conta creditada.",
        style="Normal",
    )
    add_two_col_table(
        doc,
        table_peer,
        ("Fórmula", "Quantidade de contas"),
        [
            ("1ª fórmula", "Uma conta debitada e uma conta creditada: 1 débito x 1 crédito."),
            ("2ª fórmula", "Uma conta debitada e duas ou mais contas creditadas: 1 débito x 2 ou mais créditos."),
            ("3ª fórmula", "Duas ou mais contas debitadas e uma conta creditada: 2 ou mais débitos x 1 crédito."),
            ("4ª fórmula", "Duas ou mais contas debitadas e duas ou mais contas creditadas: 2 ou mais débitos x 2 ou mais créditos."),
        ],
    )
    add_labeled_paragraph(
        doc,
        "Leitura do lançamento manual: ",
        "“a Diversos” indica mais de uma conta creditada; “Diversos”, sem a preposição, indica mais de uma conta debitada.",
    )

    add_heading(doc, "Erros e retificação de lançamentos", 2, ("retificacao_lancamentos", 105))
    doc.add_paragraph(
        "Erros comuns envolvem título da conta, valor, inversão, duplo registro, omissão e histórico incorreto. Toda retificação deve indicar, no histórico, o motivo da correção, a data e a localização do lançamento de origem.",
        style="Normal",
    )
    add_two_col_table(
        doc,
        table_peer,
        ("Técnica", "Efeito"),
        [
            ("Estorno", "Lançamento inverso que anula integralmente o registro errado; em regra, é seguido do lançamento correto."),
            ("Transferência", "Estorno parcial que desloca o valor da conta indevida para a conta correta em um único lançamento."),
            ("Complementação", "Aumenta ou reduz o valor já registrado pela diferença necessária, mantendo as mesmas contas do lançamento original."),
        ],
    )

    add_heading(doc, "Escrituração de operações típicas", 2, ("operacoes_tipicas", 106))
    add_two_col_table(
        doc,
        table_peer,
        ("Operação", "Lançamento conceitual e classificação"),
        [
            ("Integralização de capital em dinheiro", "D - Caixa / C - Capital Social. Fato modificativo aumentativo: +A e +PL."),
            ("Depósito bancário", "D - Bancos / C - Caixa. Fato permutativo: +A e -A."),
            ("Compra de mercadorias a prazo", "D - Estoques / C - Duplicatas a Pagar. Fato permutativo: +A e +P."),
            ("Aplicação financeira", "Na aplicação: D - Aplicações / C - Bancos, fato permutativo. O rendimento é apropriado como receita; o resgate inverte a transferência patrimonial."),
            ("Empréstimo bancário", "Na obtenção: D - Bancos / C - Empréstimos, fato permutativo. Os juros são despesa por competência; a quitação baixa Bancos e Empréstimos."),
            ("Venda de mercadorias", "Há pelo menos dois registros: D - Caixa ou Clientes / C - Receita de Vendas; D - CMV / C - Estoques. Com lucro, o conjunto é fato misto aumentativo."),
            ("Folha de pagamento", "Na competência: D - Despesa com Salários / C - Salários a Pagar, fato modificativo diminutivo. O pagamento posterior é permutativo."),
            ("Dívida paga com juros", "D - Obrigação e D - Juros Passivos / C - Caixa ou Bancos. Fato misto diminutivo."),
            ("Dívida paga com desconto", "D - Obrigação / C - Caixa ou Bancos e C - Desconto Obtido. Fato misto aumentativo."),
            ("Aquisição parcialmente à vista e a prazo", "D - Bem adquirido / C - Caixa e C - Financiamento a Pagar. Fato permutativo, pois não altera o PL."),
        ],
    )

    add_heading(doc, "Desconto de duplicatas", 2, ("desconto_duplicatas", 107))
    doc.add_paragraph(
        "No desconto de duplicatas, a entidade antecipa no banco o recebimento de títulos, mas permanece responsável se o cliente não pagar. Por sua essência de financiamento, Duplicatas Descontadas é passivo; Encargos Financeiros a Transcorrer é conta retificadora desse passivo.",
        style="Normal",
    )
    add_labeled_paragraph(
        doc,
        "Registro inicial: ",
        "D - Bancos pelo valor líquido; D - Encargos Financeiros a Transcorrer pelos juros antecipados; C - Duplicatas Descontadas pelo valor nominal. Os encargos são apropriados como despesa financeira pelo decorrer do tempo.",
    )
    add_two_col_table(
        doc,
        table_peer,
        ("Situação no vencimento", "Lançamento essencial"),
        [
            ("Cliente paga integralmente", "D - Duplicatas Descontadas / C - Duplicatas a Receber."),
            ("Cliente não paga", "D - Duplicatas Descontadas / C - Bancos. A duplicata a receber permanece até o tratamento da inadimplência."),
            ("Cliente paga parcialmente", "D - Duplicatas Descontadas / C - Duplicatas a Receber pela parcela paga / C - Bancos pela parcela suportada pela entidade."),
        ],
    )

    add_heading(doc, "Balancete de verificação", 2, ("balancete_verificacao", 108))
    doc.add_paragraph(
        "O balancete é demonstrativo auxiliar e não obrigatório, elaborado para fins operacionais com as contas e os saldos extraídos do Razão. Seu objetivo é verificar a igualdade decorrente do método das partidas dobradas.",
        style="Normal",
    )
    add_labeled_paragraph(
        doc,
        "Elementos mínimos: ",
        "identificação da entidade, data, abrangência, contas e grupos, indicação dos saldos devedores ou credores e soma de ambos os lados.",
    )
    add_two_col_table(
        doc,
        table_peer,
        ("Modelo", "Estrutura"),
        [
            ("2 colunas", "Saldo devedor e saldo credor."),
            ("4 colunas", "Movimentos devedor e credor, mais saldos devedor e credor."),
            ("6 colunas", "Saldo anterior D/C, movimento D/C e saldo atual D/C."),
            ("8 colunas", "Saldo anterior D/C, movimento D/C, saldos do período D/C e saldo atual D/C."),
        ],
    )
    add_labeled_paragraph(
        doc,
        "Limite do teste: ",
        "totais iguais são condição necessária, mas não provam que toda a escrituração está correta. Omissão, erro de título e determinados registros em duplicidade podem não ser detectados.",
    )
    add_labeled_paragraph(
        doc,
        "Sintético x analítico: ",
        "o balancete sintético mostra contas principais; o analítico detalha subcontas de segundo e demais graus. A coluna com o nome das contas não entra na contagem de 2, 4, 6 ou 8 colunas.",
    )

    integration = find_paragraph(
        doc,
        exact="Integração: da conta ao fato contábil",
        style_name="Heading 1",
    )
    move_new_block_before(doc, first._p, integration._p)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    if any(paragraph.text.startswith("Aula 04") for paragraph in doc.paragraphs):
        raise RuntimeError("Aula 04 is already present; refusing to duplicate it")

    subtitle = find_paragraph(doc, startswith="Concursos fiscais | Aulas 02 e 03")
    replace_text_keep_format(
        subtitle,
        "Concursos fiscais | Aulas 02 a 04 - Contas, Atos e Fatos Contábeis e Escrituração",
    )

    source_note = find_paragraph(doc, startswith="Fonte e recorte:")
    replace_text_keep_format(
        source_note,
        source_note.text.rstrip()
        + " Aula 04: 04-Escrituracao.pdf, páginas 3 a 33; a página 34 inicia as Questões Comentadas - FGV e não foi incorporada.",
    )

    toc_peer = find_paragraph(doc, exact="Aula 03 - Atos e Fatos Contábeis")
    toc_integration = find_paragraph(
        doc,
        exact="Integração: da conta ao fato contábil",
        style_name="Normal",
    )
    toc_entries = [
        ("Aula 04 - Escrituração", "aula_04"),
        ("Aspectos gerais e obrigatoriedade", "escrituracao_aspectos"),
        ("Livros de escrituração", "livros_escrituracao"),
        ("Formalidades da escrituração", "formalidades_escrituracao"),
        ("Lançamentos contábeis e fórmulas", "lancamentos_formulas"),
        ("Erros e retificação de lançamentos", "retificacao_lancamentos"),
        ("Escrituração de operações típicas", "operacoes_tipicas"),
        ("Desconto de duplicatas", "desconto_duplicatas"),
        ("Balancete de verificação", "balancete_verificacao"),
    ]
    for text, anchor in toc_entries:
        toc_integration._p.addprevious(clone_toc_paragraph(toc_peer, text, anchor))
    replace_hyperlink_text(toc_integration, "Integração: das contas à escrituração")

    bullet_peer = find_paragraph(doc, exact="Patrimoniais: representam Ativo, Passivo e Patrimônio Líquido.")
    table_peer = doc.tables[3]
    append_aula04(doc, bullet_peer, table_peer)

    integration_heading = find_paragraph(
        doc,
        exact="Integração: da conta ao fato contábil",
        style_name="Heading 1",
    )
    replace_text_keep_format(integration_heading, "Integração: das contas à escrituração")
    integration_intro = find_paragraph(doc, startswith="As contas fornecem o vocabulário;")
    replace_text_keep_format(
        integration_intro,
        "As contas fornecem o vocabulário, os fatos mostram as variações patrimoniais e a escrituração transforma cada operação em registros no Diário, no Razão e no balancete. Para resolver uma operação, siga a sequência:",
    )
    keyword_heading = find_paragraph(doc, exact="Palavras-chave")
    new_steps = [
        add_bullet(doc, bullet_peer, "Registre local e data, contas debitadas e creditadas, histórico e valor no Diário."),
        add_bullet(doc, bullet_peer, "Transfira os movimentos ao Razão e confira a igualdade dos saldos no balancete, sem presumir que a igualdade elimina todos os erros."),
    ]
    insert_new_paragraphs_before(keyword_heading, new_steps)

    keywords = find_paragraph(doc, startswith="conta • saldo")
    replace_text_keep_format(
        keywords,
        keywords.text
        + " • escrituração • Diário • Razão • Caixa • formalidades extrínsecas • formalidades intrínsecas • fórmulas de lançamento • estorno • transferência • complementação • duplicatas descontadas • balancete de verificação",
    )

    quick_heading = find_paragraph(doc, exact="Revisão rápida", style_name="Heading 1")
    new_traps = [
        add_bullet(doc, bullet_peer, "Diário x Razão x Caixa: Diário é cronológico; Razão é sistemático por conta; Caixa registra recebimentos e pagamentos.", "Diário x Razão x Caixa:"),
        add_bullet(doc, bullet_peer, "Extrínseca x intrínseca: descumprir formalidade extrínseca invalida o livro; descumprir formalidade intrínseca invalida o registro atingido.", "Extrínseca x intrínseca:"),
        add_bullet(doc, bullet_peer, "2ª x 3ª fórmula: a segunda tem um débito e vários créditos; a terceira tem vários débitos e um crédito.", "2ª x 3ª fórmula:"),
        add_bullet(doc, bullet_peer, "Balancete: igualdade entre totais devedores e credores não comprova ausência de todos os erros.", "Balancete:"),
        add_bullet(doc, bullet_peer, "Duplicatas descontadas: são passivo; encargos financeiros a transcorrer são sua conta retificadora.", "Duplicatas descontadas:"),
    ]
    insert_new_paragraphs_before(quick_heading, new_traps)

    auto_heading = find_paragraph(doc, exact="Autoavaliação", style_name="Heading 2")
    quick_items = [
        add_bullet(doc, bullet_peer, "Escrituração: registra fatos por lançamentos permanentes e documentados."),
        add_bullet(doc, bullet_peer, "Livros: Diário = cronológico; Razão = sistemático; Caixa = recebimentos e pagamentos."),
        add_bullet(doc, bullet_peer, "Formalidades: extrínsecas cuidam do livro; intrínsecas cuidam do lançamento."),
        add_bullet(doc, bullet_peer, "Fórmulas: 1ª = 1x1; 2ª = 1x2+; 3ª = 2+x1; 4ª = 2+x2+."),
        add_bullet(doc, bullet_peer, "Retificação: estorno anula; transferência corrige a conta; complementação ajusta o valor."),
        add_bullet(doc, bullet_peer, "Desconto de duplicatas: obrigação com o banco e encargos apropriados por competência."),
        add_bullet(doc, bullet_peer, "Balancete: pode ter 2, 4, 6 ou 8 colunas e não detecta todos os erros."),
    ]
    insert_new_paragraphs_before(auto_heading, quick_items)

    final_questions = [
        "Diferencie Livro Diário, Livro Razão e Livro Caixa.",
        "Quais são os elementos mínimos de um lançamento contábil?",
        "Qual é a diferença entre a 2ª e a 3ª fórmula de lançamento?",
        "Quando usar estorno, transferência e complementação?",
        "Por que Duplicatas Descontadas é passivo e como os encargos são apropriados?",
        "O que mostram os balancetes de 2, 4, 6 e 8 colunas?",
        "Por que a igualdade do balancete não garante escrituração sem erros?",
    ]
    last_existing = doc.paragraphs[-1]
    for text in final_questions:
        paragraph = add_bullet(doc, bullet_peer, text)
        last_existing._p.addnext(paragraph._p)
        last_existing = paragraph

    close = doc.add_paragraph(style="Normal")
    close.add_run("Fechamento da Aula 04: ").bold = True
    close.add_run(
        "conteúdo teórico das páginas 3 a 33 consolidado. A página 34 inicia as Questões Comentadas - FGV e ficou fora da revisão."
    )
    last_existing._p.addnext(close._p)

    doc.core_properties.subject = "Revisão cumulativa de Contabilidade para concursos fiscais - Aulas 02 a 04"
    doc.core_properties.comments = (
        "Aula 04 - Escrituração incorporada com base nas páginas 3 a 33 de 04-Escrituracao.pdf; questões comentadas excluídas."
    )
    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
