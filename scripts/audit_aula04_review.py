from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


FILE = Path(r"C:\dev\hpo-ptbr-lab\output\Contabilidade - Revisão Cumulativa - Contas, Atos e Fatos Contábeis.docx")
doc = Document(FILE)
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

required = [
    "Aula 02 - Contas",
    "Aula 03 - Atos e Fatos Contábeis",
    "Aula 04 - Escrituração",
    "Aspectos gerais e obrigatoriedade",
    "Livros de escrituração",
    "Formalidades da escrituração",
    "Lançamentos contábeis e fórmulas",
    "Erros e retificação de lançamentos",
    "Escrituração de operações típicas",
    "Desconto de duplicatas",
    "Balancete de verificação",
    "páginas 3 a 33",
    "página 34 inicia",
    "Fechamento da Aula 04",
]
missing = [value for value in required if value not in text]
assert not missing, f"Missing required content: {missing}"

for forbidden in [
    "==43512d==",
    "01996352342",
    "Esterfânia Araujo Barbosa Farias",
    "(FGV/PC-RN/2021)",
    "Em 01/07/2020, a Cia. Alfa contratou um seguro",
]:
    assert forbidden not in text, f"Forbidden source/question content found: {forbidden}"

headings = [
    paragraph.text
    for paragraph in doc.paragraphs
    if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
]
idx_02 = headings.index("Aula 02 - Contas")
idx_03 = headings.index("Aula 03 - Atos e Fatos Contábeis")
idx_04 = headings.index("Aula 04 - Escrituração")
idx_integration = headings.index("Integração: das contas à escrituração")
idx_review = headings.index("Revisão rápida")
assert idx_02 < idx_03 < idx_04 < idx_integration < idx_review

assert text.count("Aula 04 - Escrituração") == 2, "Aula 04 must appear once in TOC and once in body"
assert len(doc.tables) == 17, f"Unexpected table count: {len(doc.tables)}"

for index, table in enumerate(doc.tables, start=1):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    assert tr_pr.find(qn("w:tblHeader")) is not None, f"Table {index} lacks repeating header"

with ZipFile(FILE) as archive:
    xml = archive.read("word/document.xml")
root = etree.fromstring(xml)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
bookmarks = root.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
anchors = root.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
assert len(bookmarks) == len(set(bookmarks)), "Duplicate bookmark names"
missing_anchors = sorted(set(anchors) - set(bookmarks))
assert not missing_anchors, f"TOC anchors without bookmarks: {missing_anchors}"

new_bookmarks = {
    "aula_04",
    "escrituracao_aspectos",
    "livros_escrituracao",
    "formalidades_escrituracao",
    "lancamentos_formulas",
    "retificacao_lancamentos",
    "operacoes_tipicas",
    "desconto_duplicatas",
    "balancete_verificacao",
}
assert new_bookmarks <= set(bookmarks), f"Missing bookmarks: {new_bookmarks - set(bookmarks)}"
assert new_bookmarks <= set(anchors), f"Missing TOC links: {new_bookmarks - set(anchors)}"

print("CUSTOM AUDIT: OK")
print(f"Paragraphs: {len(doc.paragraphs)} | Tables: {len(doc.tables)}")
print(f"Bookmarks: {len(bookmarks)} | Internal links: {len(anchors)}")
print(f"File size: {FILE.stat().st_size} bytes")
